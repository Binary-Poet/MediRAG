"""多格式解析测试：内存构造样本，不读真实磁盘文件。"""
import io
import json

import pytest

from app.ingestion.parsers import SUPPORTED_EXTS, ParseError, parse_document


def _docx_bytes() -> bytes:
    from docx import Document as Docx
    d = Docx()
    d.add_paragraph("四君子汤由人参、白术、茯苓、炙甘草组成。")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["方剂", "组成"])
    ws.append(["四君子汤", "人参、白术"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pptx_bytes() -> bytes:
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "脾气虚证候要点"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _pdf_bytes() -> bytes:
    """最小可用 PDF（手写单页对象结构，含文本流）。"""
    content = b"BT /F1 12 Tf 72 720 Td (Sijunzi Tang) Tj ET"
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += b"trailer\n<< /Size " + str(len(objs) + 1).encode() + b" /Root 1 0 R >>\n"
    out += b"startxref\n" + str(xref_pos).encode() + b"\n%%EOF\n"
    return bytes(out)


@pytest.mark.parametrize("ext", ["txt", "md", "csv", "tsv", "json", "xml", "yaml", "yml", "log", "html"])
def test_plain_text_family(ext):
    text = parse_document("人参大补元气。".encode("utf-8"), ext)
    assert "人参" in text


def test_json_is_pretty_printed():
    raw = json.dumps({"方剂": "四君子汤"}, ensure_ascii=False).encode("utf-8")
    assert "四君子汤" in parse_document(raw, "json")


def test_docx_extraction():
    assert "四君子汤" in parse_document(_docx_bytes(), "docx")


def test_xlsx_extraction():
    text = parse_document(_xlsx_bytes(), "xlsx")
    assert "方剂" in text and "四君子汤" in text


def test_pptx_extraction():
    assert "脾气虚" in parse_document(_pptx_bytes(), "pptx")


def test_pdf_extraction():
    assert "Sijunzi" in parse_document(_pdf_bytes(), "pdf")


def test_unsupported_ext_raises():
    with pytest.raises(ParseError):
        parse_document(b"x", "exe")


def test_binary_doc_raises_parse_error():
    with pytest.raises(ParseError):
        parse_document(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1\x00\x01\x02", "doc")


def test_supported_exts_cover_spec_list():
    spec = {"pdf", "doc", "docx", "rtf", "ppt", "pptx", "xls", "xlsx", "csv", "tsv",
            "txt", "md", "html", "json", "xml", "yaml", "log"}
    assert spec <= SUPPORTED_EXTS